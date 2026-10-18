from __future__ import annotations

import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt


def add_markdownish_text(doc: Document, text: str) -> None:
    """Small Markdown-to-DOCX converter for headings, bullets, and paragraphs."""
    in_code = False
    for raw_line in text.splitlines():
        line = raw_line.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if not line.strip():
            continue
        if in_code:
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            continue
        match = re.match(r"^(#{1,3})\s+(.*)$", line)
        if match:
            level = min(len(match.group(1)), 3)
            doc.add_heading(match.group(2).strip(), level=level)
        elif re.match(r"^\s*[-*]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\s*\d+[.)]\s+", line):
            doc.add_paragraph(re.sub(r"^\s*\d+[.)]\s+", "", line), style="List Number")
        else:
            # Lightweight bold support.
            p = doc.add_paragraph()
            parts = re.split(r"(\*\*.*?\*\*)", line)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    p.add_run(part[2:-2]).bold = True
                else:
                    p.add_run(part)


def compile_docx(title: str, brief: str, chapters: list[dict], output_path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    doc.add_heading(title, 0)
    doc.add_paragraph("Draft manuscript — requires author review")
    doc.add_heading("Book Brief", level=1)
    add_markdownish_text(doc, brief)

    doc.add_page_break()
    doc.add_heading("Contents", level=1)
    for item in chapters:
        doc.add_paragraph(f"Chapter {item['number']}: {item['title']}")

    for item in chapters:
        doc.add_page_break()
        doc.add_heading(f"Chapter {item['number']}: {item['title']}", level=1)
        add_markdownish_text(doc, item["draft"])
        doc.add_heading("Research dossier", level=2)
        add_markdownish_text(doc, item["research"])
        doc.add_heading("Review report", level=2)
        add_markdownish_text(doc, item["review_text"])
        doc.add_heading("Diagram source", level=2)
        doc.add_paragraph(f"Mermaid source saved separately: diagrams/chapter-{item['number']:02d}.mmd")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
